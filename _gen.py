# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

DOCX_SRC = r"C:\Users\18780\Desktop\1234567\水墨神州_项目计划书_技术部分.docx"
DOCX_OUT = r"C:\Users\18780\Desktop\1234567\水墨神州_项目计划书_含图片.docx"
IMG_DIR = r"C:\Users\18780\Desktop\1234567\doc_images"

# Each tuple: (image filename, keyword to match in chapter heading)
IMAGE_MAP = [
    ("图1_项目技术架构总览.png", "项目技术概述"),
    ("图2_四层架构模型.png", "系统总体架构"),
    ("图3_地图沙盘系统.png", "核心技术模块"),
    ("图4_三维建筑展示系统.png", "三维古建筑"),
    ("图5_AI技术集成架构.png", "人工智能技术"),
    ("图6_游戏化教育系统.png", "游戏化教育"),
    ("图7_数据库ER图.png", "数据管理"),
    ("图8_水墨风格UI体系.png", "用户界面与交互"),
    ("图9_创新点总览.png", "系统创新点"),
    ("图10_技术发展路线图.png", "技术发展路线"),
]

def insert_images():
    if not os.path.exists(DOCX_SRC):
        print(f"Error: docx not found: {DOCX_SRC}")
        sys.exit(1)

    doc = Document(DOCX_SRC)

    # Build a list of (paragraph_index, image_path) for insertion
    insert_plan = []
    for img_file, keyword in IMAGE_MAP:
        img_path = os.path.join(IMG_DIR, img_file)
        if not os.path.exists(img_path):
            print(f"Warning: image not found: {img_path}")
            continue

        for i, para in enumerate(doc.paragraphs):
            if keyword in para.text:
                insert_plan.append((i, img_path))
                print(f"  Matched '{keyword}' at paragraph {i}: {para.text[:40]}")
                break
        else:
            print(f"  Warning: no paragraph found for keyword '{keyword}'")

    # Insert images in reverse order so paragraph indices stay valid
    insert_plan.sort(key=lambda x: x[0], reverse=True)

    for para_idx, img_path in insert_plan:
        # We can't directly insert at an index with python-docx,
        # so we add the picture at the end then move it.
        # Instead, use the paragraph's underlying XML element to insert after.
        from docx.oxml.ns import qn
        from lxml import etree

        target_para = doc.paragraphs[para_idx]

        # Create a new paragraph for the image
        new_para = doc.add_paragraph()  # temporarily at end
        new_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = new_para.add_run()
        run.add_picture(img_path, width=Cm(14))

        # Move the new paragraph element right after the target paragraph
        target_elem = target_para._element
        new_elem = new_para._element
        target_elem.addnext(new_elem)

        print(f"  Inserted {os.path.basename(img_path)} after paragraph {para_idx}")

    doc.save(DOCX_OUT)
    print(f"\nDone! Saved to: {DOCX_OUT}")

if __name__ == "__main__":
    print("Inserting images into document...")
    insert_images()
